import pandas as pd
from pathlib import Path
import magic
from openpyxl import load_workbook
from time import perf_counter
from datetime import datetime

COM=[]
Food_wastage=[]
Delivery=[]
Veg_utilisation=[]
Utilisation_cost=[]
Electricity_utilisation=[]
Gas_utilisation=[]
Water_utilsation=[]
briquets_utilisation=[]
Deisel_utilisation=[]
production=[]

output=Path(r"Output path")

logfile=Path(r"Input path")
def clean_date(value):

    if pd.isna(value):
        return pd.NaT

    value = str(value).strip()

    formats = [
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M:%S",
        "%d-%b-%Y",
        "%d/%m/%Y",
        "%d-%m-%Y"
    ]

    for fmt in formats:
        try:
            return pd.to_datetime(value, format=fmt)
        except (ValueError, TypeError):
            continue

    return pd.NaT
def log_process(mtime,log):
    df=pd.read_csv(logfile)
    df.loc[len(df)]={"relative_path":log,"Modified_time":mtime}
    print("Logged file Sucessfully")
    df.to_csv(logfile,index=False)

def check_processed(log,mtime):
    if logfile.exists():
        print("Log file found")
        df=pd.read_csv(logfile)
        if(log,mtime) in df.values:
            return True
    else:
        
        df=pd.DataFrame(columns=["relative_path","Modified_time"])
        df.to_csv(logfile,index=False)
        print("Created Dataframe sucessfylly")
        print(df.head(5))
        return False

input_file=Path(r"Input File)
try:

    if input_file.exists():
        print("Found path")
        if input_file.is_dir():
            print(input_file.name,":Is Dir")
            file=list(input_file.iterdir())
            print("No of files:",len(file))
            for files in file:
                print(type(files))
            
                if files.is_dir():
                    print(files.name,":is a folder")
                    subfile=list(files.iterdir())
                    print("no of subfiles:",len(subfile))
                    for subfiles in subfile:
                        if subfiles.stem.lower()=="navigation":
                            print("Skipping navigation")
                            continue
                        else:

                            print(subfiles.name)
                            relative_path=subfiles.relative_to(input_file)
                            modified_time=datetime.fromtimestamp(subfiles.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
                            print("Relative path:",relative_path)
                            print("Modified time:",modified_time)
                            
                            if check_processed(relative_path,modified_time):
                                print("Already processed")
                                continue
                            else:
                                print("Starting load workbook:",subfiles.name)
                                workbook_start=perf_counter()
                                excel=load_workbook(subfiles)
                                workbook_end=perf_counter()
                                print("Workbook load time:",workbook_end-workbook_start)
                                
                                for sheetname in excel.sheetnames:
                                    if sheetname.lower().strip().replace(" ","").replace("_","") in ["fixedcosts","costfixed","fixedcost"]:
                                        print("Skipping:",sheetname)
                                        continue
                                    else:
                                        
                                        start=perf_counter()
                                        print("Reading:",sheetname)
                                        read_start=perf_counter()
                                        df=pd.read_excel(subfiles,sheet_name=sheetname,header=1)
                                        read_end=perf_counter()
                                        print(f"Read time:{read_end-read_start:.2f}sec")
                                        clean_start=perf_counter()
                                        df=df.replace(["X","x"],pd.NA)
                                        df=df.dropna(axis=1,thresh=7)
                                        
                                        print(df.dtypes)
                                        df["Kitchen"]=subfiles.parent.stem
                                        df=df.rename(columns={"date":"Date"})
                                        print("Renamed Columns")
                                        print(sheetname,":",df.columns)
                                        if "Date" in df.columns:
                                            df["Date"]=df["Date"].astype(str).replace(".","/")
                                            print("Replaced values")
                                            
                                            df["Date"]=df["Date"].ffill()
                                            df["Date"]=df["Date"].apply(clean_date)
                                            print("Filled Dates")
                                        else:
                                            print("Date column not found")
                                        
                                            clean_end=perf_counter()
                                            print(f"Cleaning time:{clean_end-clean_start:.2f}sec")
                                            print(df.head(5))
                                            read_end=perf_counter()
                                            end=perf_counter()
                                            print(" total Time taken:",end-start)
                                        
                                        if sheetname.lower().strip().replace("_","").replace(" ","") in ["com"]:
                                            COM.append(df)
                                        elif sheetname.lower().strip().replace("_","") in ["production","productionmis"]:
                                            production.append(df)
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["vegutilisation"]:
                                            Veg_utilisation.append(df)
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["utilisationcost","costutilisation"]:
                                            Utilisation_cost.append(df)
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["gasutilisation"]:
                                            Gas_utilisation.append(df)
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["electricityutilisation"]:
                                            Electricity_utilisation.append(df)
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["waterutilisation"]:
                                                Water_utilsation.append(df)
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["briquetsutilisation","briquietsutilisation","briqueitsutilisation"]:
                                            briquets_utilisation.append(df)
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["deiselutilisation","dieselutilsation"]:
                                            Deisel_utilisation.append(df)
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["delivery","canteendelivery"]:
                                            Delivery.append(df)
                                            print("Delivery files",len(Delivery))
                                        elif sheetname.lower().strip().replace("_","").replace(" ","") in ["foodwastage","food wastage","foodwastagereport"]:
                                            print("foodwastage file")
                                            Food_wastage.append(df)
                                        log_process(modified_time,relative_path)
    if len(COM)<=0:
        print("No files to cocatenate in com:",len(COM))
    else:
        final_com=pd.concat(COM,ignore_index=True)
        print("COM:",final_com.head(5))
    if(len(Food_wastage)<=0):
        print("Empty dataframe no files to concat in food wastage",len(Food_wastage))
    else:
        final_foodwastage=pd.concat(Food_wastage,ignore_index=True)
        print(final_foodwastage.head(5))
    if len(Delivery)<=0:
            print("no files to concatentae in Delivery:",len(Delivery))
    else:
        final_delivery=pd.concat(Delivery,ignore_index=True)
    if len(Water_utilsation)<=0:
        print("No sufficent files to concatenate Water utilisation",len(Water_utilsation))
    else:
        final_water_utilisation=pd.concat(Water_utilsation,ignore_index=True)
    if len(Veg_utilisation)<=0:
        print("Insufficient files to concatenate veg utilisation",len(Veg_utilisation))
    else:
        final_vegutilisation=pd.concat(Veg_utilisation,ignore_index=True)
    if len(Deisel_utilisation)<=0:
        print("Insufficient files to concatenate deisel Utilisation",len(Deisel_utilisation))
    else:
        final_deiselutilisation=pd.concat(Deisel_utilisation,ignore_index=True)
    if len(Gas_utilisation)<=0:
        print("Insufficient files to concatenate in Gas utilisation",len(Gas_utilisation))
    else:
        final_gasutilsation=pd.concat(Gas_utilisation,ignore_index=True)
    if len(Electricity_utilisation)<=0:
        print("Insufficient files to concatenate electricty Utilisation:",len(Electricity_utilisation))
    else:
        final_Electricity_utilisation=pd.concat(Electricity_utilisation,ignore_index=True)
    if len(briquets_utilisation)<=0:
        print("Insufficient files to concatenate briquiets Utilisation:",len(briquets_utilisation))
    else:
        final_briquets_utilisation=pd.concat(briquets_utilisation,ignore_index=True)
    if len(Utilisation_cost)<=0:
        print("Insufficient files to concatenate utilisation Cost:",len(Utilisation_cost))
    else:
        final_utilisation_cost=pd.concat(Utilisation_cost,ignore_index=True)
    if(len(production)<=0):
        print("Insufficient Files to Concatenate in Production:",len(production))
    else:
        final_production=pd.concat(production,ignore_index=True)
    


    #==== export to csv======#
    final_delivery.to_csv(output/"Delivery.csv",index=False)
    print("Expoprted Delivery file:",output.parent.name)
    final_foodwastage.to_csv(output/"Food Wastage.csv",index=False)
    print("Exported food wastage file:",output.parent.name)
    final_briquets_utilisation.to_csv(output/"Briquiets utilisation.csv",index=False)
    print("Exported briquets utilisation to:",output.parent.name)
    final_deiselutilisation.to_csv(output/"Deisel utilisation.csv",index=False)
    print("Exported deisel utilisation to:",output.parent.name)
    final_Electricity_utilisation.to_csv(output/"Electricity utilisation.csv",index=False)
    print("Exported electricity utilisation",output.parent.name)
    final_gasutilsation.to_csv(output/"Gas utilisation.csv",index=False)
    print("Exported gas utilisation:",output.parent.name)
    final_water_utilisation.to_csv(output/"Water utilisation.csv",index=False)
    print("Exported water utilisation to:",output.parent.name)
    final_com.to_csv(output/"COM.csv",index=False)
    print("Exported COM to:",output.parent.name)
    final_vegutilisation.to_csv(output/"vegutilisation.csv",index=False)
    print("Exported Veg utilisationto:",output.parent.name)
    final_utilisation_cost.to_csv(output/"utilisationcost.csv",index=False)
    print("Exported utilisation cost to:",output.parent.name)
    final_production.to_csv(output/"production.csv",index=False)
    print("Exported production File to:",output.parent.name)

    print("all exports sucessfull")

    


        
    


except Exception as e:
    print("Folder:",subfiles.parent.stem)
    print("file:",subfiles.name)
    print("Sheet:",sheetname)
    print("Error type:",type(e).__name__)
    print("Error:",e)
    raise
   



else:
    print("Given path not valid")
